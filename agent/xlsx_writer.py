"""生成用例写回 Excel：**始终在副本上操作，原文件只读**。

为什么不做原地修改
------------------
``docs/cases/ERP系统测试用例.xlsx`` 是业务评审的真源，里面有人工修正的
预期结果和手工执行状态。LLM 生成的用例直接写进去会污染历史数据、破坏
人工修改，也无法回溯是哪次生成改的。

所以流程固定为：

1. 复制原文件到 ``docs/cases/generated/ERP系统测试用例_<timestamp>.xlsx``；
2. 只在副本的目标 sheet **追加**行，不删不改已有行；
3. 原文件的内容与 mtime 都不变。
"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from agent.parser.extractor import EXCEL_COLUMNS

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DOCS_CASES = PROJECT_ROOT / "docs" / "cases"
GENERATED_DIR = DOCS_CASES / "generated"
SOURCE_WORKBOOK = DOCS_CASES / "ERP系统测试用例.xlsx"

# LLM 可能输出英文优先级，统一成 Excel 的中文取值
PRIORITY_LABELS = {"high": "高", "medium": "中", "low": "低"}


def copy_workbook(
    source: Path = SOURCE_WORKBOOK,
    out_dir: Path = GENERATED_DIR,
) -> Path:
    """复制源工作簿，返回副本路径。"""
    if not source.exists():
        raise FileNotFoundError(f"源用例文件不存在：{source}")
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    dest = out_dir / f"{source.stem}_{stamp}{source.suffix}"
    shutil.copy2(source, dest)
    return dest


def _text(value, default: str = "") -> str:
    if value is None:
        return default
    if isinstance(value, (list, tuple)):
        return "\n".join(str(item) for item in value)
    return str(value)


def _case_to_row(case: dict) -> list[str]:
    """把 LLM 输出的一条 JSON 用例映射成 12 列。"""
    priority = _text(case.get("priority"), "高").strip()
    priority = PRIORITY_LABELS.get(priority.lower(), priority)

    return [
        _text(case.get("id")),
        priority,
        _text(case.get("module")),
        _text(case.get("purpose")),
        _text(case.get("preconditions")),
        _text(case.get("steps")),
        _text(case.get("expected_result")),
        _text(case.get("assertions")),
        _text(case.get("manual_status"), "待执行"),
        _text(case.get("automation_status"), "待生成"),
        _text(case.get("tester"), "AI"),
        _text(case.get("ai_mark"), "AI生成"),
    ]


def append_cases(copy_path: Path, cases: list[dict], *, sheet: str | None = None) -> int:
    """在副本的目标 sheet 追加用例，返回追加条数。"""
    if not cases:
        return 0

    workbook = load_workbook(copy_path)
    try:
        worksheet = workbook[sheet] if sheet and sheet in workbook.sheetnames else workbook.active

        headers = [cell.value for cell in worksheet[1]]
        if headers[: len(EXCEL_COLUMNS)] == list(EXCEL_COLUMNS):
            column_index = {name: index for index, name in enumerate(EXCEL_COLUMNS)}
        else:
            # 表头顺序被调整过：按表头名映射，避免写错列
            column_index = {name: index for index, name in enumerate(headers) if name}

        added = 0
        for case in cases:
            width = len(headers) if headers else len(EXCEL_COLUMNS)
            values = [None] * width
            for column, name in enumerate(EXCEL_COLUMNS):
                target = column_index.get(name)
                if target is not None:
                    values[target] = _case_to_row(case)[column]
            worksheet.append(values)
            added += 1

        workbook.save(copy_path)
        return added
    finally:
        workbook.close()


def write_generated_cases(
    cases: list[dict],
    *,
    source: Path = SOURCE_WORKBOOK,
    out_dir: Path = GENERATED_DIR,
    sheet: str | None = None,
) -> tuple[Path, int]:
    """复制源文件 → 在副本追加 → 返回 ``(副本路径, 追加条数)``。"""
    if not cases:
        raise ValueError("没有可写入的用例")
    dest = copy_workbook(source, out_dir)
    added = append_cases(dest, cases, sheet=sheet)
    return dest, added
