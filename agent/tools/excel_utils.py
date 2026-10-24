"""Excel 用例读取与步骤解析工具。"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from agent.Agents.Base_Agent import PROJECT_ROOT

XLSX_DEFAULT = PROJECT_ROOT / "docs" / "cases" / "ERP系统测试用例.xlsx"

REQUIRED_COLUMNS = (
    "用例编号",
    "优先级",
    "功能模块",
    "测试目的",
    "前置条件",
    "测试步骤",
    "预期结果",
    "断言校验",
    "自动化执行情况",
)


def parse_steps(raw: str) -> list[str]:
    """``1.点击新增\n2.选择供应商`` -> ``['点击新增', '选择供应商']``。"""
    steps: list[str] = []
    for line in (raw or "").splitlines():
        text = line.strip()
        if not text:
            continue
        match = re.match(r"^\d+[.、]\s*(.+)$", text)
        steps.append(match.group(1).strip() if match else text)
    return steps


def load_pending_cases(xlsx: Path, sheet: str) -> list[dict]:
    """读取 sheet，只返回 ``自动化执行情况 == 待执行`` 的用例。"""
    if not xlsx.exists():
        raise SystemExit(f"[截断] Excel 不存在：{xlsx}")

    workbook = load_workbook(xlsx, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise SystemExit(f"[截断] Excel 没有 sheet：{sheet}")
        worksheet = workbook[sheet]
        headers = [str(cell.value or "").strip() for cell in worksheet[1]]
        index = {name: position for position, name in enumerate(headers)}

        missing = [name for name in REQUIRED_COLUMNS if name not in index]
        if missing:
            raise SystemExit(f"[截断] sheet {sheet} 缺少列：{missing}")

        cases: list[dict] = []
        last_module = ""
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[index["用例编号"]]:
                continue
            automation = str(row[index["自动化执行情况"]] or "").strip()
            if automation != "待执行":
                continue

            module = str(row[index["功能模块"]] or "").strip() or last_module
            last_module = module

            cases.append(
                {
                    "id": str(row[index["用例编号"]]).strip(),
                    "priority": str(row[index["优先级"]] or "").strip(),
                    "module": module,
                    "purpose": str(row[index["测试目的"]] or "").strip(),
                    "precondition": str(row[index["前置条件"]] or "").strip(),
                    "steps": parse_steps(str(row[index["测试步骤"]] or "")),
                    "expected": str(row[index["预期结果"]] or "").strip(),
                    "assertions": str(row[index["断言校验"]] or "").strip(),
                }
            )
        return cases
    finally:
        workbook.close()
