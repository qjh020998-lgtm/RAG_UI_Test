"""读取 Excel 工作簿，转换为 TestSpec。

只负责 openpyxl 读取和行列提取；列名映射、模块继承等由 extractor 处理。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from agent.models import TestSpec
from agent.parser.extractor import EXCEL_COLUMNS, extract_from_excel_rows


def _read_sheet(ws: Worksheet, *, sheet: str, workbook_path: str) -> TestSpec:
    headers = [ws.cell(1, col).value for col in range(1, len(EXCEL_COLUMNS) + 1)]
    if not any(headers):
        raise ValueError(f"{sheet} 缺少表头")

    rows = [
        [ws.cell(row, col).value for col in range(1, len(EXCEL_COLUMNS) + 1)]
        for row in range(2, ws.max_row + 1)
    ]
    return extract_from_excel_rows(
        rows,
        sheet=sheet,
        workbook_path=workbook_path,
        title=f"{sheet} — ERP 测试用例",
    )


def read_sheet(path: str | Path, sheet: str) -> TestSpec:
    """读取单个 sheet。"""
    workbook_path = str(path)
    wb = load_workbook(workbook_path, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise KeyError(f"sheet 不存在: {sheet}")
        return _read_sheet(wb[sheet], sheet=sheet, workbook_path=workbook_path)
    finally:
        wb.close()


def read_workbook(path: str | Path) -> dict[str, TestSpec]:
    """读取全部非空 sheet。"""
    workbook_path = str(path)
    wb = load_workbook(workbook_path, data_only=True)
    try:
        result: dict[str, TestSpec] = {}
        for ws in wb.worksheets:
            if ws.max_row < 2:
                continue
            result[ws.title] = _read_sheet(
                ws,
                sheet=ws.title,
                workbook_path=workbook_path,
            )
        return result
    finally:
        wb.close()


__all__ = ["read_sheet", "read_workbook"]
