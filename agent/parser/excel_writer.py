"""把 TestSpec 写回 Excel，保留模板样式、合并单元格和列宽。"""

from __future__ import annotations

from copy import copy
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from agent.models import TestCase, TestSpec
from agent.parser.excel_reader import read_sheet
from agent.parser.extractor import EXCEL_COLUMNS


def _case_row(case: TestCase) -> list[str]:
    priority = case.priority.value if hasattr(case.priority, "value") else str(case.priority)
    return [
        case.id,
        priority,
        case.module,
        case.purpose,
        "\n".join(case.preconditions),
        "\n".join(case.steps),
        case.expected_result,
        case.assertions,
        case.manual_status,
        case.automation_status,
        case.tester,
        case.ai_mark,
    ]


def _write_header(ws: Worksheet) -> None:
    for col, value in enumerate(EXCEL_COLUMNS, 1):
        cell = ws.cell(1, col, value)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _clear_c_merges(ws: Worksheet) -> None:
    for rng in list(ws.merged_cells.ranges):
        if rng.min_col == 3 and rng.max_col == 3:
            ws.unmerge_cells(str(rng))


def _merge_modules(ws: Worksheet, cases: list[TestCase]) -> None:
    if not cases:
        return
    start = 2
    for index in range(1, len(cases)):
        if (cases[index].module or "") == (cases[index - 1].module or ""):
            continue
        end = index + 1
        if end > start:
            ws.merge_cells(start_row=start, start_column=3, end_row=end, end_column=3)
            ws.cell(start, 3).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        start = index + 2
    end = len(cases) + 1
    if end > start:
        ws.merge_cells(start_row=start, start_column=3, end_row=end, end_column=3)
        ws.cell(start, 3).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _replace_cases(ws: Worksheet, cases: Iterable[TestCase]) -> None:
    cases = list(cases)
    template_styles = []
    row_height = 42
    if ws.max_row >= 2:
        row_height = ws.row_dimensions[2].height or row_height
        template_styles = [copy(ws.cell(2, col)._style) for col in range(1, 13)]

    _clear_c_merges(ws)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for row_index, case in enumerate(cases, start=2):
        for col_index, value in enumerate(_case_row(case), start=1):
            cell = ws.cell(row_index, col_index, value)
            if template_styles:
                cell._style = copy(template_styles[col_index - 1])
            if col_index in {5, 6, 7, 8}:
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row_index].height = row_height

    _merge_modules(ws, cases)
    ws.auto_filter.ref = f"A1:L{max(1, len(cases) + 1)}"


def write_spec(
    path: str | Path,
    spec: TestSpec,
    *,
    sheet: str | None = None,
    mode: str = "replace",
    template_sheet: str | None = None,
) -> Path:
    """写入 TestSpec。

    mode:
    - ``replace``：覆盖目标 sheet 的数据区；
    - ``append``：读取已有数据后追加。
    新 sheet 可复制 ``template_sheet`` 的样式。
    """
    if mode not in {"replace", "append"}:
        raise ValueError("mode 只能是 replace 或 append")

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    sheet = sheet or spec.sheet or "Sheet1"
    exists = target.exists()
    wb = load_workbook(target) if exists else Workbook()
    sheet_existed = sheet in wb.sheetnames

    if not sheet_existed:
        if template_sheet and template_sheet in wb.sheetnames:
            ws = wb.copy_worksheet(wb[template_sheet])
            ws.title = sheet
        else:
            ws = wb.active if len(wb.sheetnames) == 1 and wb.active.title == "Sheet" else wb.create_sheet()
            ws.title = sheet
            _write_header(ws)
    else:
        ws = wb[sheet]

    cases = list(spec.test_cases)
    if mode == "append" and exists and sheet_existed:
        existing = read_sheet(target, sheet)
        existing_ids = {case.id for case in existing.test_cases}
        cases = existing.test_cases + [
            case for case in cases if case.id not in existing_ids
        ]

    _replace_cases(ws, cases)
    wb.save(target)
    return target


def update_statuses(
    path: str | Path,
    sheet: str,
    updates: Mapping[str, Mapping[str, str]],
) -> int:
    """按用例编号更新 I/J 状态列，返回实际更新条数。"""
    wb = load_workbook(path)
    if sheet not in wb.sheetnames:
        raise KeyError(f"sheet 不存在: {sheet}")
    ws = wb[sheet]
    rows: dict[str, int] = {}
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if value is None:
            continue
        # 单元格可能是 ArrayFormula/CellRichText 等不可哈希对象，
        # 不能直接作为 dict key；统一转成字符串。
        case_id = str(value).strip()
        if case_id:
            rows[case_id] = row
    updated = 0
    for case_id, values in updates.items():
        row = rows.get(case_id)
        if row is None:
            continue
        if "manual" in values:
            ws.cell(row, 9).value = values["manual"]
        if "automation" in values:
            ws.cell(row, 10).value = values["automation"]
        updated += 1
    wb.save(path)
    return updated


__all__ = ["write_spec", "update_statuses"]
